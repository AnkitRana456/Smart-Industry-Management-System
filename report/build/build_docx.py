import os
import re
import sys

sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_report as B
import content_front as F
from content_ch678 import BIB

OUT = B.ROOT + r'\Smart_Industrial_Management_System_Final_Report.docx'
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.8)
sec.top_margin = Inches(0.6)
sec.bottom_margin = Inches(0.8)

SPLIT = re.compile(r'(\*\*.+?\*\*|`.+?`|(?<![\w*])_.+?_(?![\w*]))')


def runs(p, text, size=None, bold=False):
    for part in SPLIT.split(text):
        if not part:
            continue
        b, i, c = bold, False, False
        if part.startswith('**'):
            part, b = part[2:-2], True
        elif part.startswith('`'):
            part, c = part[1:-1], True
        elif part.startswith('_') and part.endswith('_') and len(part) > 2:
            part, i = part[1:-1], True
        r = p.add_run(part)
        r.bold, r.italic = b, i
        if c:
            r.font.name = 'Consolas'
        if size:
            r.font.size = Pt(size)


def para(text, size=None, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    runs(p, text, size)


def head(text, size, upper=False):
    p = doc.add_paragraph()
    if upper:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(text.upper() if upper else text)
    r.bold = True
    r.font.size = Pt(size)


def table(headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        runs(c.paragraphs[0], str(h), size, True)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ''
            runs(cells[i].paragraphs[0], str(v), size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(6.9 * w / 100)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(p, text, 11, True)


def picture(path, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(min(width, 6.9)))


def sec_pg(section, fmt):
    el = section._sectPr.find(qn('w:pgNumType'))
    if el is None:
        el = OxmlElement('w:pgNumType')
        section._sectPr.append(el)
    el.set(qn('w:fmt'), fmt)
    el.set(qn('w:start'), '1')


def footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.text = ''
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), alignment=2)
    r = p.add_run(B.DEPT + '\t')
    r.font.size = Pt(10)
    r = p.add_run()
    r.font.size = Pt(10)
    for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), kind)
        else:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = txt
        r._r.append(e)


def new_section(fmt):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    footer(s)
    sec_pg(s, fmt)


def pbreak():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def render(b, title=None):
    t = b['t']
    if t == 'h1':
        head(b['text'], 16, upper=True)
    elif t == 'h2':
        head(b['text'], 14)
    elif t == 'h3':
        head(b['text'], 12)
    elif t == 'p':
        para(b['text'])
    elif t in ('ul', 'ol'):
        for x in b['items']:
            para(x, style='List Bullet' if t == 'ul' else 'List Number')
    elif t == 'table':
        caption(b['caption'])
        table(b['headers'], b['rows'], b['widths'], b['size'])
    elif t == 'fig':
        picture(b['img'], b['width'])
        caption(b['caption'])
    elif t == 'figtext':
        picture(b['img'], min(b['width'], 3.4))
        caption(b['caption'])
        for x in b['blocks']:
            render(x)
    elif t == 'raw':
        h = b['html']
        if 'class="sig"' in h:
            table(['', '', ''], [[f'({n})\n{e}' for n, e in F.MEMBERS]], None, 12)
        elif 'toc-row' in h:
            for k, text, lab in B.toc:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(0 if k == 'h1' else 14)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), alignment=2)
                r = p.add_run(re.sub(r'[`*]', '', text) + '\t' + lab)
                r.font.size = Pt(10.5)
                r.bold = k == 'h1'
        elif 'class="lst"' in h:
            noun, rows = ('Figure', B.lof) if title == 'List of Figures' else ('Table', B.lot)
            table([noun + ' No.', 'Title', 'Page No.'],
                  [[c.split(':')[0].replace(noun + ' ', ''), c.split(':', 1)[1].strip(), p]
                   for c, p in (r.split('|') for r in rows)], [14, 74, 12], 10.5)
        elif 'class="ref"' in h:
            for r in BIB:
                para(r, size=12)


cover = [(F.TITLE, 24, True), (F.SUBTITLE, 16, True), ('(ES-451 : Minor Project)', 24, True), ('', 12, False),
         ('Submitted in partial fulfillment of the requirement for the award of the degree of', 12, False),
         ('Bachelor of Technology', 16, True), ('in', 12, False), ('Computer Science & Technology', 14, True),
         ('', 12, False), ('Submitted By', 12, False)] + [(f'{n} ({e})', 12, True) for n, e in F.MEMBERS] + \
        [('', 12, False), ('Under the supervision of', 12, False), ('Dr. Garima', 12, True), ('Professor (CST)', 12, True)]
for txt, sz, bd in cover:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(txt)
    r.bold, r.font.size = bd, Pt(sz)
picture(F.LOGO, 1.3)
for txt, sz in (('COMPUTER SCIENCE AND TECHNOLOGY', 16), ('MAHARAJA AGRASEN INSTITUTE OF TECHNOLOGY', 16),
                ('ROHINI SEC-22, NEW DELHI – 110086', 14), ('Nov 2026', 16)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(txt)
    r.bold, r.font.size = True, Pt(sz)

new_section('lowerRoman')
for i in range(1, B.FRONT):
    pg = B.pages[i]
    title = pg[0]['text'] if pg[0]['t'] == 'h1' else None
    if i > 1:
        pbreak()
    for b in pg:
        render(b, title)
new_section('decimal')
for j, pg in enumerate(B.body_pages):
    if j:
        pbreak()
    for b in pg:
        render(b)
doc.save(OUT)
print('saved', OUT)
